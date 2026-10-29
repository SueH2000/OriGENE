#!/usr/bin/env python3
"""Test legacy OriGENE fold models on fully unseen genes from another patient.

This script reuses existing fold checkpoints.  It does not retrain a model and
does not apply every checkpoint to every target gene.

The historical ``Predictions_*_Fold*_genes.npy`` files are row manifests.  A
gene may occur twice because the legacy experiment contained two observations
for that gene.  Therefore, being present in a held-out manifest is not enough:
the same gene may also occur in another fold and hence in the training rows for
the current checkpoint.

For checkpoint/fold k this script constructs:

``training_gene_set[k]``
    Unique gene IDs occurring in the other ``n_folds - 1`` row folds.

``eligible_target_gene_set[k]``
    Unique gene IDs occurring in held-out fold k but absent from
    ``training_gene_set[k]``.

Checkpoint k predicts only the corresponding eligible genes from the target
patient.  Thus the evaluated subset satisfies, by construction::

    source_training_gene_set[k] ∩ evaluated_target_gene_set[k] = ∅

The fold-specific results are reported separately and then concatenated into a
pooled target-patient result.  Predictions from models that trained on a gene
are never used for that gene.

The preferred proof source is the legacy per-fold gene artifact expected by
``Average performance of 5 fold.ipynb``::

    Predictions_{source}_Fold{fold}_genes.npy

If those artifacts are missing, deterministic fold reconstruction is available
behind ``--reconstruct-folds``.  ``--force-reconstruct-folds`` deliberately
ignores any saved manifests (for example, manifests from a different combined-
patient experiment) and reconstructs the source-only folds.  Reconstruction
can enforce a clean evaluation assignment, but it cannot independently prove
the training provenance or fold numbering of an old checkpoint.

Examples
--------
Audit legacy artifacts without loading TensorFlow models::

    python cross_patient_gene_disjoint_reuse_weights.py --audit-only

Run the Lung I source models against all configured target patients::

    python cross_patient_gene_disjoint_reuse_weights.py

Use a different explicit held-out-gene filename template::

    python cross_patient_gene_disjoint_reuse_weights.py \
      --heldout-template 'Predictions_{source}_Fold{fold}_genes.npy'

Reconstruct folds only when the exact original source arrays and order are
known to be unchanged::

    python cross_patient_gene_disjoint_reuse_weights.py --reconstruct-folds

Ignore unrelated saved manifests and run the no-retraining, reconstructed
analysis::

    python cross_patient_gene_disjoint_reuse_weights.py \
      --force-reconstruct-folds

The output directory includes a formatted Excel workbook, publication-ready
PNG/PDF figures, fold-specific and pooled CSV results, a gene eligibility
manifest, and a JSON provenance record.  The output proves the set exclusion
performed by this script.  It cannot independently prove that a legacy
checkpoint corresponds to a historical row manifest; that requires the
original training log/indices or exact prediction reproduction.
"""

from __future__ import annotations

import argparse
import gc
import hashlib
import json
import os
import sys
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, Iterable, List, Mapping, Sequence, Tuple

import numpy as np
import pandas as pd
from sklearn.metrics import (
    accuracy_score,
    average_precision_score,
    confusion_matrix,
    f1_score,
    precision_recall_curve,
    precision_score,
    recall_score,
    roc_auc_score,
    roc_curve,
)
from sklearn.model_selection import StratifiedKFold


DEFAULT_TARGETS = (
    "Lung_II",
    "Liver_I",
    "Liver_II",
    "Colorectal_II",
    "Colorectal_III",
)


@dataclass(frozen=True)
class DatasetSpec:
    root_kind: str
    normal_name: str
    cancer_name: str


@dataclass(frozen=True)
class FoldEligibility:
    """Gene-ID sets implied by one legacy row-level fold assignment."""

    heldout_rows: np.ndarray
    heldout_genes: np.ndarray
    training_genes: np.ndarray
    eligible_genes: np.ndarray
    excluded_seen_genes: np.ndarray


DATASET_SPECS: Mapping[str, DatasetSpec] = {
    "Lung_I": DatasetSpec("base", "Lung_Normal", "Lung_Cancer"),
    "Lung_II": DatasetSpec("base", "Lung_Normal_I", "Lung_Cancer_I"),
    "Liver_I": DatasetSpec("base", "Liver_Normal", "Liver_Cancer"),
    "Liver_II": DatasetSpec("base", "Liver_Normal_I", "Liver_Cancer_I"),
    "Colorectal_II": DatasetSpec(
        "new", "Colorectal_Normal_II", "Colorectal_Cancer_II"
    ),
    "Colorectal_III": DatasetSpec(
        "new", "Colorectal_Normal_III", "Colorectal_Cancer_III"
    ),
}


class AuditError(RuntimeError):
    """Raised when a condition needed for gene-disjoint reuse is not proven."""


def canonical_gene(value: object) -> str:
    """Return a conservative canonical gene ID without changing its biology."""
    gene = str(value).strip()
    if not gene:
        raise AuditError("Encountered an empty gene identifier.")
    return gene


def canonical_genes(values: Iterable[object]) -> np.ndarray:
    return np.asarray([canonical_gene(value) for value in values], dtype=str)


def expand(path: str | Path) -> Path:
    return Path(os.path.expandvars(os.path.expanduser(str(path)))).resolve()


def sha256_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while True:
            block = handle.read(chunk_size)
            if not block:
                break
            digest.update(block)
    return digest.hexdigest()


def sha256_strings(values: Iterable[str]) -> str:
    digest = hashlib.sha256()
    for value in values:
        digest.update(value.encode("utf-8"))
        digest.update(b"\0")
    return digest.hexdigest()


def require_unique_genes(genes: np.ndarray, context: str) -> None:
    duplicated = pd.Index(genes)[pd.Index(genes).duplicated()].unique().tolist()
    if duplicated:
        preview = ", ".join(map(str, duplicated[:10]))
        raise AuditError(
            f"{context} contains duplicated gene IDs ({len(duplicated)} unique "
            f"duplicates). First examples: {preview}"
        )


def array_paths(root: Path, dataset_name: str) -> Tuple[Path, Path, Path]:
    return (
        root / f"Final_sequences_40kbp_{dataset_name}.npy",
        root / f"Final_targets_40kbp_{dataset_name}.npy",
        root / f"Final_genes_40kbp_{dataset_name}.npy",
    )


def load_metadata(root: Path, dataset_name: str) -> Tuple[np.ndarray, np.ndarray]:
    _, target_path, gene_path = array_paths(root, dataset_name)
    for path in (target_path, gene_path):
        if not path.is_file():
            raise FileNotFoundError(f"Missing dataset file: {path}")

    labels = np.asarray(np.load(target_path, allow_pickle=True)).reshape(-1).astype(int)
    genes = canonical_genes(np.load(gene_path, allow_pickle=True).reshape(-1))
    if len(labels) != len(genes):
        raise AuditError(
            f"{dataset_name}: {len(labels)} labels but {len(genes)} gene IDs."
        )
    require_unique_genes(genes, dataset_name)
    return labels, genes


def load_triplet(
    root: Path, dataset_name: str
) -> Tuple[np.ndarray, np.ndarray, np.ndarray]:
    sequence_path, target_path, gene_path = array_paths(root, dataset_name)
    for path in (sequence_path, target_path, gene_path):
        if not path.is_file():
            raise FileNotFoundError(f"Missing dataset file: {path}")

    x = np.asarray(np.load(sequence_path, allow_pickle=True))
    y = np.asarray(np.load(target_path, allow_pickle=True)).reshape(-1).astype(int)
    genes = canonical_genes(np.load(gene_path, allow_pickle=True).reshape(-1))

    if x.ndim == 2:
        x = np.expand_dims(x, axis=-1)
    if x.ndim != 3:
        raise AuditError(
            f"{dataset_name}: expected a 2D or 3D sequence array, got {x.shape}."
        )
    if not (len(x) == len(y) == len(genes)):
        raise AuditError(
            f"{dataset_name}: inconsistent lengths X={len(x)}, y={len(y)}, "
            f"genes={len(genes)}."
        )
    require_unique_genes(genes, dataset_name)
    return x, y, genes


def index_by_gene(genes: np.ndarray, context: str) -> Dict[str, int]:
    require_unique_genes(genes, context)
    return {gene: index for index, gene in enumerate(genes)}


def align_pair_metadata(
    normal_labels: np.ndarray,
    normal_genes: np.ndarray,
    cancer_labels: np.ndarray,
    cancer_genes: np.ndarray,
    patient: str,
) -> Tuple[np.ndarray, np.ndarray]:
    normal_set = set(normal_genes)
    cancer_set = set(cancer_genes)
    if normal_set != cancer_set:
        only_normal = sorted(normal_set - cancer_set)
        only_cancer = sorted(cancer_set - normal_set)
        raise AuditError(
            f"{patient}: normal/cancer gene sets differ. "
            f"normal-only={only_normal[:10]}, cancer-only={only_cancer[:10]}."
        )

    normal_pos = index_by_gene(normal_genes, f"{patient} normal")
    normal_order = np.asarray([normal_pos[gene] for gene in cancer_genes], dtype=int)
    if not np.array_equal(normal_labels[normal_order], cancer_labels):
        raise AuditError(f"{patient}: normal/cancer class labels disagree by gene ID.")
    return cancer_labels.copy(), cancer_genes.copy()


def align_pair_arrays(
    x_normal: np.ndarray,
    y_normal: np.ndarray,
    g_normal: np.ndarray,
    x_cancer: np.ndarray,
    y_cancer: np.ndarray,
    g_cancer: np.ndarray,
    patient: str,
) -> Tuple[np.ndarray, np.ndarray, np.ndarray]:
    labels, genes = align_pair_metadata(
        y_normal, g_normal, y_cancer, g_cancer, patient
    )
    normal_pos = index_by_gene(g_normal, f"{patient} normal")
    normal_order = np.asarray([normal_pos[gene] for gene in genes], dtype=int)
    x_normal = x_normal[normal_order]

    if x_normal.shape[:2] != x_cancer.shape[:2]:
        raise AuditError(
            f"{patient}: aligned normal/cancer shapes differ: "
            f"{x_normal.shape} versus {x_cancer.shape}."
        )
    if x_normal.shape[-1] != 1 or x_cancer.shape[-1] != 1:
        raise AuditError(
            f"{patient}: expected one channel per normal/cancer array before "
            f"pairing; got {x_normal.shape[-1]} and {x_cancer.shape[-1]}."
        )

    x_paired = np.concatenate([x_normal, x_cancer], axis=-1)
    return x_paired, labels, genes


def dataset_root(spec: DatasetSpec, sequence_root: Path, new_root: Path) -> Path:
    if spec.root_kind == "base":
        return sequence_root
    if spec.root_kind == "new":
        return new_root
    raise AuditError(f"Unknown dataset root kind: {spec.root_kind}")


def load_patient_metadata(
    patient: str, sequence_root: Path, new_root: Path
) -> Tuple[np.ndarray, np.ndarray]:
    if patient not in DATASET_SPECS:
        raise AuditError(
            f"Unknown patient '{patient}'. Choices: {sorted(DATASET_SPECS)}"
        )
    spec = DATASET_SPECS[patient]
    root = dataset_root(spec, sequence_root, new_root)
    y_normal, g_normal = load_metadata(root, spec.normal_name)
    y_cancer, g_cancer = load_metadata(root, spec.cancer_name)
    return align_pair_metadata(y_normal, g_normal, y_cancer, g_cancer, patient)


def load_patient_arrays(
    patient: str, sequence_root: Path, new_root: Path
) -> Tuple[np.ndarray, np.ndarray, np.ndarray]:
    if patient not in DATASET_SPECS:
        raise AuditError(
            f"Unknown patient '{patient}'. Choices: {sorted(DATASET_SPECS)}"
        )
    spec = DATASET_SPECS[patient]
    root = dataset_root(spec, sequence_root, new_root)
    x_normal, y_normal, g_normal = load_triplet(root, spec.normal_name)
    x_cancer, y_cancer, g_cancer = load_triplet(root, spec.cancer_name)
    return align_pair_arrays(
        x_normal,
        y_normal,
        g_normal,
        x_cancer,
        y_cancer,
        g_cancer,
        patient,
    )


def load_patient_gene_subset(
    patient: str,
    selected_genes: Sequence[str],
    sequence_root: Path,
    new_root: Path,
) -> Tuple[np.ndarray, np.ndarray, np.ndarray]:
    """Load and pair only selected genes to avoid materializing the full tensor."""
    if patient not in DATASET_SPECS:
        raise AuditError(
            f"Unknown patient '{patient}'. Choices: {sorted(DATASET_SPECS)}"
        )
    chosen = canonical_genes(selected_genes)
    require_unique_genes(chosen, f"{patient} selected target genes")
    if len(chosen) == 0:
        raise AuditError(f"{patient}: no target genes were selected for evaluation.")

    spec = DATASET_SPECS[patient]
    root = dataset_root(spec, sequence_root, new_root)
    normal_labels, normal_genes = load_metadata(root, spec.normal_name)
    cancer_labels, cancer_genes = load_metadata(root, spec.cancer_name)
    labels, _ = align_pair_metadata(
        normal_labels,
        normal_genes,
        cancer_labels,
        cancer_genes,
        patient,
    )
    normal_position = index_by_gene(normal_genes, f"{patient} normal")
    cancer_position = index_by_gene(cancer_genes, f"{patient} cancer")
    missing = [
        gene
        for gene in chosen
        if gene not in normal_position or gene not in cancer_position
    ]
    if missing:
        raise AuditError(
            f"{patient}: {len(missing)} selected genes are absent from the paired "
            f"target arrays; examples={missing[:10]}."
        )

    normal_sequence_path, _, _ = array_paths(root, spec.normal_name)
    cancer_sequence_path, _, _ = array_paths(root, spec.cancer_name)
    for path in (normal_sequence_path, cancer_sequence_path):
        if not path.is_file():
            raise FileNotFoundError(f"Missing dataset file: {path}")
    x_normal = np.load(normal_sequence_path, allow_pickle=True, mmap_mode="r")
    x_cancer = np.load(cancer_sequence_path, allow_pickle=True, mmap_mode="r")
    if x_normal.ndim == 2:
        x_normal = np.expand_dims(x_normal, axis=-1)
    if x_cancer.ndim == 2:
        x_cancer = np.expand_dims(x_cancer, axis=-1)
    if x_normal.ndim != 3 or x_cancer.ndim != 3:
        raise AuditError(
            f"{patient}: expected 2D or 3D sequence arrays, got "
            f"{x_normal.shape} and {x_cancer.shape}."
        )
    if len(x_normal) != len(normal_genes) or len(x_cancer) != len(cancer_genes):
        raise AuditError(
            f"{patient}: sequence and gene-array lengths do not agree."
        )
    if x_normal.shape[1:] != x_cancer.shape[1:]:
        raise AuditError(
            f"{patient}: normal/cancer sequence shapes differ: "
            f"{x_normal.shape} versus {x_cancer.shape}."
        )
    if x_normal.shape[-1] != 1 or x_cancer.shape[-1] != 1:
        raise AuditError(
            f"{patient}: expected one channel per normal/cancer array before "
            f"pairing; got {x_normal.shape[-1]} and {x_cancer.shape[-1]}."
        )

    normal_indices = np.asarray([normal_position[gene] for gene in chosen], dtype=int)
    cancer_indices = np.asarray([cancer_position[gene] for gene in chosen], dtype=int)
    cancer_label_map = dict(zip(cancer_genes, labels))
    selected_labels = np.asarray(
        [cancer_label_map[gene] for gene in chosen], dtype=int
    )
    paired = np.concatenate(
        [x_normal[normal_indices], x_cancer[cancer_indices]], axis=-1
    )
    return paired, selected_labels, chosen


def format_template(template: str, source: str, fold: int) -> str:
    try:
        return template.format(source=source, fold=fold)
    except KeyError as error:
        raise AuditError(
            f"Template '{template}' contains unsupported field {error}. "
            "Only {source} and {fold} are allowed."
        ) from error


def resolve_under(base_dir: Path, rendered: str) -> Path:
    expanded_text = os.path.expandvars(os.path.expanduser(rendered))
    candidate = Path(expanded_text)
    if not candidate.is_absolute():
        candidate = base_dir / candidate
    candidate = candidate.resolve()
    return candidate


def resolve_model_paths(
    model_dir: Path, template: str, source: str, n_folds: int
) -> Dict[int, Path]:
    paths: Dict[int, Path] = {}
    for fold in range(1, n_folds + 1):
        rendered = format_template(template, source, fold)
        path = resolve_under(model_dir, rendered)
        if not path.is_file():
            raise FileNotFoundError(f"Missing fold-{fold} checkpoint: {path}")
        paths[fold] = path

    if len(set(paths.values())) != n_folds:
        raise AuditError(
            "Checkpoint template resolves multiple folds to the same path. "
            "This reproduces the legacy stale-variable bug and is forbidden."
        )

    hashes = {fold: sha256_file(path) for fold, path in paths.items()}
    if len(set(hashes.values())) != n_folds:
        duplicates: Dict[str, List[int]] = {}
        for fold, digest in hashes.items():
            duplicates.setdefault(digest, []).append(fold)
        repeated = {key: value for key, value in duplicates.items() if len(value) > 1}
        raise AuditError(
            "Two or more fold checkpoints have identical SHA-256 content, so "
            f"distinct fold provenance cannot be established: {repeated}"
        )
    return paths


def resolve_heldout_path(
    heldout_dir: Path,
    template: str | None,
    source: str,
    fold: int,
) -> Path:
    if template:
        rendered = format_template(template, source, fold)
        path = resolve_under(heldout_dir, rendered)
        if not path.is_file():
            raise FileNotFoundError(f"Missing fold-{fold} held-out gene file: {path}")
        return path

    exact_names = (
        f"Predictions_{source}_Fold{fold}_genes.npy",
        f"OriGENE_{source}_Fold{fold}_genes.npy",
    )
    exact = [heldout_dir / name for name in exact_names if (heldout_dir / name).is_file()]
    if len(exact) == 1:
        return exact[0]
    if len(exact) > 1:
        raise AuditError(
            f"Fold {fold}: multiple exact held-out gene files found: {exact}. "
            "Pass --heldout-template explicitly."
        )

    globbed = sorted(heldout_dir.glob(f"*Fold{fold}*genes.npy"))
    globbed = [path for path in globbed if "mut" not in path.name.lower()]
    if len(globbed) == 1:
        return globbed[0]
    if not globbed:
        raise FileNotFoundError(
            f"Fold {fold}: no held-out gene file found in {heldout_dir}. "
            "Pass --heldout-template or, only if the exact training arrays are "
            "unchanged, use --reconstruct-folds."
        )
    raise AuditError(
        f"Fold {fold}: ambiguous held-out gene files: {globbed}. "
        "Pass --heldout-template explicitly."
    )


def load_heldout_folds(
    heldout_dir: Path,
    template: str | None,
    source: str,
    source_genes: np.ndarray,
    source_labels: np.ndarray,
    n_folds: int,
    reconstruct: bool,
    force_reconstruct: bool,
    random_state: int,
) -> Tuple[Dict[int, np.ndarray], Dict[int, Path], str]:
    fold_genes: Dict[int, np.ndarray] = {}
    fold_paths: Dict[int, Path] = {}

    if force_reconstruct:
        splitter = StratifiedKFold(
            n_splits=n_folds, shuffle=True, random_state=random_state
        )
        for fold, (_, heldout_index) in enumerate(
            splitter.split(np.zeros(len(source_labels)), source_labels), start=1
        ):
            fold_genes[fold] = source_genes[heldout_index].copy()
        proof_source = (
            "forced_source_only_reconstruction_"
            "unproven_checkpoint_fold_provenance"
        )
    else:
        try:
            for fold in range(1, n_folds + 1):
                path = resolve_heldout_path(heldout_dir, template, source, fold)
                genes = canonical_genes(np.load(path, allow_pickle=True).reshape(-1))
                fold_genes[fold] = genes
                fold_paths[fold] = path
            proof_source = "saved_row_fold_gene_manifests"
        except FileNotFoundError:
            if not reconstruct:
                raise
            splitter = StratifiedKFold(
                n_splits=n_folds, shuffle=True, random_state=random_state
            )
            fold_genes = {}
            for fold, (_, heldout_index) in enumerate(
                splitter.split(np.zeros(len(source_labels)), source_labels), start=1
            ):
                fold_genes[fold] = source_genes[heldout_index].copy()
            fold_paths = {}
            proof_source = (
                "deterministically_reconstructed_"
                "unproven_checkpoint_fold_provenance"
            )

    return fold_genes, fold_paths, proof_source


def derive_fold_eligibility(
    fold_rows: Mapping[int, np.ndarray], source_genes: np.ndarray
) -> Tuple[Dict[int, FoldEligibility], Dict[str, object]]:
    """Derive genes completely absent from each checkpoint's training folds.

    The other folds are the training rows for checkpoint ``k``.  Consequently,
    a held-out gene is eligible only when it does not occur in any other fold.
    Duplicate occurrences inside the same held-out fold are allowed and collapse
    to one target-patient gene-level prediction.
    """
    source_set = set(source_genes)
    fold_sets = {fold: set(rows.tolist()) for fold, rows in fold_rows.items()}
    unexpected = {
        fold: sorted(genes - source_set)
        for fold, genes in fold_sets.items()
        if genes - source_set
    }
    if unexpected:
        raise AuditError(
            f"Row-fold manifests contain genes outside the source universe: {unexpected}"
        )

    manifest_union = set().union(*fold_sets.values())
    missing = sorted(source_set - manifest_union)
    if missing:
        raise AuditError(
            "Row-fold manifests do not cover the source gene universe. "
            f"covered={len(manifest_union)}, source={len(source_set)}, "
            f"missing examples={missing[:10]}."
        )

    eligibility: Dict[int, FoldEligibility] = {}
    pairwise_overlap: Dict[str, int] = {}
    for left in sorted(fold_sets):
        for right in sorted(fold_sets):
            if right > left:
                pairwise_overlap[f"{left}-{right}"] = len(
                    fold_sets[left] & fold_sets[right]
                )

    for fold in sorted(fold_sets):
        heldout_set = fold_sets[fold]
        training_set = set().union(
            *(genes for other, genes in fold_sets.items() if other != fold)
        )
        eligible_set = heldout_set - training_set
        excluded_set = heldout_set & training_set
        eligibility[fold] = FoldEligibility(
            heldout_rows=fold_rows[fold].copy(),
            heldout_genes=np.asarray(sorted(heldout_set), dtype=str),
            training_genes=np.asarray(sorted(training_set), dtype=str),
            eligible_genes=np.asarray(sorted(eligible_set), dtype=str),
            excluded_seen_genes=np.asarray(sorted(excluded_set), dtype=str),
        )

    eligible_sets = {
        fold: set(details.eligible_genes) for fold, details in eligibility.items()
    }
    eligible_pairwise_overlap: Dict[str, int] = {}
    for left in sorted(eligible_sets):
        for right in sorted(eligible_sets):
            if right > left:
                overlap = eligible_sets[left] & eligible_sets[right]
                eligible_pairwise_overlap[f"{left}-{right}"] = len(overlap)
                if overlap:
                    raise AuditError(
                        "Internal error: fully unseen eligible genes were assigned "
                        f"to folds {left} and {right}: {sorted(overlap)[:10]}"
                    )

    eligible_union = set().union(*eligible_sets.values())
    occurrence_counts = Counter(
        gene for rows in fold_rows.values() for gene in rows.tolist()
    )
    occurrence_distribution = Counter(occurrence_counts.values())
    return eligibility, {
        "source_unique_gene_count": len(source_set),
        "manifest_total_row_count": int(sum(len(rows) for rows in fold_rows.values())),
        "manifest_unique_gene_count": len(manifest_union),
        "manifest_gene_occurrence_distribution": {
            str(copies): int(count)
            for copies, count in sorted(occurrence_distribution.items())
        },
        "fold_row_counts": {
            str(fold): int(len(rows)) for fold, rows in fold_rows.items()
        },
        "fold_unique_heldout_gene_counts": {
            str(fold): int(len(details.heldout_genes))
            for fold, details in eligibility.items()
        },
        "fold_source_training_gene_counts": {
            str(fold): int(len(details.training_genes))
            for fold, details in eligibility.items()
        },
        "fold_fully_unseen_eligible_gene_counts": {
            str(fold): int(len(details.eligible_genes))
            for fold, details in eligibility.items()
        },
        "fold_excluded_seen_gene_counts": {
            str(fold): int(len(details.excluded_seen_genes))
            for fold, details in eligibility.items()
        },
        "pairwise_heldout_unique_gene_overlap": pairwise_overlap,
        "pairwise_eligible_gene_overlap": eligible_pairwise_overlap,
        "eligible_unique_gene_count": len(eligible_union),
        "eligible_source_coverage_fraction": float(
            len(eligible_union) / len(source_set)
        ),
        "manifest_covers_source_universe": manifest_union == source_set,
        "eligible_gene_sets_disjoint": True,
    }


def target_gene_audit(
    source_genes: np.ndarray,
    source_labels: np.ndarray,
    target_genes: np.ndarray,
    target_labels: np.ndarray,
    source: str,
    target: str,
) -> Dict[str, object]:
    source_set = set(source_genes)
    target_set = set(target_genes)
    if source_set != target_set:
        only_source = sorted(source_set - target_set)
        only_target = sorted(target_set - source_set)
        raise AuditError(
            f"{source} -> {target}: source/target gene universes differ. "
            f"source-only={len(only_source)} {only_source[:10]}; "
            f"target-only={len(only_target)} {only_target[:10]}."
        )

    source_label_map = dict(zip(source_genes, source_labels))
    target_label_map = dict(zip(target_genes, target_labels))
    disagreements = [
        gene
        for gene in source_set
        if int(source_label_map[gene]) != int(target_label_map[gene])
    ]
    if disagreements:
        raise AuditError(
            f"{source} -> {target}: labels disagree for {len(disagreements)} "
            f"genes; examples: {sorted(disagreements)[:10]}."
        )

    return {
        "target_gene_count": len(target_set),
        "source_target_gene_sets_equal": True,
        "source_target_labels_equal": True,
        "source_target_gene_overlap_count": len(source_set),
    }


def evaluated_target_subset_audit(
    source_genes: np.ndarray,
    source_labels: np.ndarray,
    evaluated_genes: np.ndarray,
    evaluated_labels: np.ndarray,
    expected_genes: Iterable[str],
    source: str,
    target: str,
) -> Dict[str, object]:
    """Validate a deliberately filtered target subset against its expectation."""
    require_unique_genes(evaluated_genes, f"{source} -> {target} evaluated subset")
    if len(evaluated_genes) != len(evaluated_labels):
        raise AuditError(
            f"{source} -> {target}: evaluated subset has {len(evaluated_genes)} "
            f"genes but {len(evaluated_labels)} labels."
        )

    source_set = set(source_genes)
    evaluated_set = set(evaluated_genes)
    expected_set = {canonical_gene(gene) for gene in expected_genes}
    if evaluated_set != expected_set:
        missing = sorted(expected_set - evaluated_set)
        unexpected = sorted(evaluated_set - expected_set)
        raise AuditError(
            f"{source} -> {target}: loaded evaluated subset differs from the "
            f"expected fully unseen set. missing={len(missing)} {missing[:10]}; "
            f"unexpected={len(unexpected)} {unexpected[:10]}."
        )
    if not evaluated_set.issubset(source_set):
        unexpected = sorted(evaluated_set - source_set)
        raise AuditError(
            f"{source} -> {target}: evaluated genes outside the source universe: "
            f"{unexpected[:10]}."
        )

    source_label_map = dict(zip(source_genes, source_labels))
    disagreements = sorted(
        gene
        for gene, label in zip(evaluated_genes, evaluated_labels)
        if int(source_label_map[gene]) != int(label)
    )
    if disagreements:
        raise AuditError(
            f"{source} -> {target}: labels disagree for "
            f"{len(disagreements)} evaluated genes; examples: "
            f"{disagreements[:10]}."
        )

    return {
        "evaluated_gene_count": len(evaluated_set),
        "evaluated_set_matches_expected_eligible_set": True,
        "evaluated_genes_are_in_source_universe": True,
        "evaluated_labels_match_source": True,
    }


def load_epi_model(path: Path):
    try:
        import keras

        return keras.models.load_model(path, compile=False, safe_mode=False)
    except ImportError:
        try:
            from tensorflow.keras.models import load_model
        except ImportError as error:
            raise RuntimeError(
                "Keras/TensorFlow is required for prediction. Run --audit-only "
                "without it, or activate the OriGENE model environment before "
                "full evaluation."
            ) from error
        return load_model(path, compile=False)


def validate_model_input(model, x_test: np.ndarray, path: Path) -> None:
    input_shape = model.input_shape
    if isinstance(input_shape, list):
        raise AuditError(
            f"{path.name} has multiple inputs {input_shape}. This script is for "
            "the epigenetic-only two-channel model and will not silently zero a "
            "mutation branch."
        )
    expected = tuple(input_shape[1:])
    observed = tuple(x_test.shape[1:])
    if len(expected) != len(observed):
        raise AuditError(
            f"{path.name}: model input {expected} does not match target {observed}."
        )
    for expected_dim, observed_dim in zip(expected, observed):
        if expected_dim is not None and int(expected_dim) != int(observed_dim):
            raise AuditError(
                f"{path.name}: model input {expected} does not match target {observed}."
            )


def clear_model(model) -> None:
    del model
    try:
        import keras

        keras.backend.clear_session()
    except ImportError:
        try:
            from tensorflow.keras import backend as keras_backend

            keras_backend.clear_session()
        except ImportError:
            pass
    gc.collect()


def classification_metrics(
    labels: np.ndarray, scores: np.ndarray, threshold: float
) -> Dict[str, float | int]:
    predicted = (scores >= threshold).astype(int)
    matrix = confusion_matrix(labels, predicted, labels=[0, 1])
    tn, fp, fn, tp = matrix.ravel()
    specificity = float(tn / (tn + fp)) if (tn + fp) else float("nan")
    has_both_classes = len(np.unique(labels)) == 2
    return {
        "N": int(len(labels)),
        "N_negative": int(np.sum(labels == 0)),
        "N_positive": int(np.sum(labels == 1)),
        "AUROC": (
            float(roc_auc_score(labels, scores))
            if has_both_classes
            else float("nan")
        ),
        "AUPRC": (
            float(average_precision_score(labels, scores))
            if has_both_classes
            else float("nan")
        ),
        "accuracy": float(accuracy_score(labels, predicted)),
        "precision": float(precision_score(labels, predicted, zero_division=0)),
        "recall": float(recall_score(labels, predicted, zero_division=0)),
        "specificity": specificity,
        "f1": float(f1_score(labels, predicted, zero_division=0)),
        "threshold": float(threshold),
        "TN": int(tn),
        "FP": int(fp),
        "FN": int(fn),
        "TP": int(tp),
    }


def bootstrap_intervals(
    labels: np.ndarray,
    scores: np.ndarray,
    n_bootstraps: int,
    seed: int,
) -> Dict[str, float | int]:
    if n_bootstraps <= 0:
        return {}
    rng = np.random.default_rng(seed)
    aurocs: List[float] = []
    auprcs: List[float] = []
    indices = np.arange(len(labels))
    for _ in range(n_bootstraps):
        sample = rng.choice(indices, size=len(indices), replace=True)
        sampled_labels = labels[sample]
        if len(np.unique(sampled_labels)) != 2:
            continue
        sampled_scores = scores[sample]
        aurocs.append(float(roc_auc_score(sampled_labels, sampled_scores)))
        auprcs.append(float(average_precision_score(sampled_labels, sampled_scores)))

    if not aurocs or not auprcs:
        raise AuditError("Bootstrap produced no valid two-class samples.")
    return {
        "bootstrap_requested": int(n_bootstraps),
        "bootstrap_valid": int(len(aurocs)),
        "AUROC_CI95_low": float(np.quantile(aurocs, 0.025)),
        "AUROC_CI95_high": float(np.quantile(aurocs, 0.975)),
        "AUPRC_CI95_low": float(np.quantile(auprcs, 0.025)),
        "AUPRC_CI95_high": float(np.quantile(auprcs, 0.975)),
    }


def make_plot(predictions: pd.DataFrame, output_dir: Path, source: str) -> Dict[str, Path]:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    plot_order = [target for target in DEFAULT_TARGETS if target in set(predictions.target)]
    plot_order += sorted(set(predictions.target) - set(plot_order))
    colors = ["cornflowerblue", "darkgreen", "limegreen", "darkred", "orange"]

    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 9,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "axes.titleweight": "bold",
        }
    )
    fig, axes = plt.subplots(1, 2, figsize=(10.5, 4.4))
    fig.patch.set_facecolor("white")
    axes[0].plot(
        [0, 1], [0, 1], color="#94A3B8", linewidth=1.0, linestyle="--",
        label="Random (0.50)"
    )

    for index, target in enumerate(plot_order):
        subset = predictions[predictions.target == target]
        labels = subset.label.to_numpy(dtype=int)
        scores = subset.score.to_numpy(dtype=float)
        color = colors[index % len(colors)]

        fpr, tpr, _ = roc_curve(labels, scores)
        precision, recall, _ = precision_recall_curve(labels, scores)
        auroc = roc_auc_score(labels, scores)
        auprc = average_precision_score(labels, scores)
        display = target.replace("_", " ")
        axes[0].plot(
            fpr, tpr, color=color, linewidth=2.0,
            label=f"{display} ({auroc:.3f})"
        )
        axes[1].plot(
            recall,
            precision,
            color=color,
            linewidth=2.0,
            label=f"{display} ({auprc:.3f})",
        )

    axes[0].set(
        xlim=(0, 1), ylim=(0, 1.02), xlabel="False positive rate",
        ylabel="True positive rate", title="ROC curves"
    )
    axes[1].set(
        xlim=(0, 1), ylim=(0, 1.02), xlabel="Recall", ylabel="Precision",
        title="Precision–recall curves"
    )
    for axis in axes:
        axis.grid(color="#E2E8F0", linewidth=0.7, alpha=0.8)
        axis.legend(fontsize=7.8, frameon=False, loc="lower right")
    fig.suptitle(
        f"Cross-patient performance from {source.replace('_', ' ')} checkpoints",
        fontsize=12,
        fontweight="bold",
        y=1.01,
    )
    fig.tight_layout()

    stem = output_dir / f"{source}_gene_disjoint_cross_patient_ROC_AUPRC"
    png_path = stem.with_suffix(".png")
    pdf_path = stem.with_suffix(".pdf")
    fig.savefig(png_path, dpi=300, bbox_inches="tight", facecolor="white")
    fig.savefig(pdf_path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return {"roc_pr_png": png_path, "roc_pr_pdf": pdf_path}


def build_gene_leakage_audit_table(
    eligibility: Mapping[int, FoldEligibility],
    source_genes: np.ndarray,
    target_audits: Mapping[str, Mapping[str, object]],
    proof_source: str,
) -> pd.DataFrame:
    """Create a reviewer-readable audit of the actually evaluated subset."""
    rows: List[Dict[str, object]] = []
    source_set = set(source_genes)
    manifest_union = set().union(
        *(set(details.heldout_genes) for details in eligibility.values())
    )
    eligible_union = set().union(
        *(set(details.eligible_genes) for details in eligibility.values())
    )

    def add(
        scope: str,
        check: str,
        expected: object,
        observed: object,
        status: str,
        interpretation: str,
    ) -> None:
        rows.append(
            {
                "scope": scope,
                "check": check,
                "expected": expected,
                "observed": observed,
                "status": status,
                "interpretation": interpretation,
            }
        )

    add(
        "Global",
        "Unique source gene IDs",
        len(source_genes),
        len(source_set),
        "PASS" if len(source_genes) == len(source_set) else "FAIL",
        "Every source gene has one stable identifier.",
    )
    add(
        "Global",
        "Row-fold manifests cover source genes",
        len(source_set),
        len(manifest_union),
        "PASS" if manifest_union == source_set else "FAIL",
        "Every source gene occurs in at least one saved row-fold manifest.",
    )
    add(
        "Global",
        "Fully unseen target-test genes retained",
        "Informational",
        f"{len(eligible_union)} / {len(source_set)}",
        "INFO",
        (
            "Only genes absent from the corresponding checkpoint's other four "
            "row folds are evaluated. Other genes are excluded, not relabeled."
        ),
    )

    for fold in sorted(eligibility):
        details = eligibility[fold]
        training_set = set(details.training_genes)
        test_set = set(details.eligible_genes)
        overlap = training_set & test_set
        add(
            f"Fold {fold}",
            "Source-training vs evaluated target-test gene overlap",
            0,
            len(overlap),
            "PASS" if not overlap else "FAIL",
            (
                f"Training unique genes={len(training_set)}; saved held-out "
                f"rows={len(details.heldout_rows)}; held-out unique genes="
                f"{len(details.heldout_genes)}; fully unseen evaluated genes="
                f"{len(test_set)}; held-out genes excluded because they also "
                f"occurred in training folds={len(details.excluded_seen_genes)}."
            ),
        )
        add(
            f"Fold {fold}",
            "Held-out genes excluded as previously seen",
            "Informational",
            len(details.excluded_seen_genes),
            "INFO",
            "These genes are not used for target-patient performance estimation.",
        )

    for target, audit in target_audits.items():
        if "pooled_prediction_count" in audit:
            predicted = int(audit["pooled_prediction_count"])
            unique = int(audit["unique_predicted_gene_count"])
            duplicates = predicted - unique
            expected_count = int(audit["eligible_gene_count"])
            status = "PASS" if predicted == expected_count and duplicates == 0 else "FAIL"
            observed: object = (
                f"{unique} unique / {predicted} rows; duplicates={duplicates}"
            )
            expected: object = (
                f"{expected_count} unique / {expected_count} rows; duplicates=0"
            )
        else:
            status = "NOT_RUN"
            observed = "Prediction stage not run"
            expected = f"{len(eligible_union)} eligible unique genes"
        add(
            target,
            "Each eligible target gene predicted exactly once",
            expected,
            observed,
            status,
            "Ineligible genes are excluded; predictions are not averaged across checkpoints.",
        )

    add(
        "Provenance",
        "Saved row-fold manifests available",
        "Saved manifests or exact original split indices",
        proof_source,
        "PASS" if proof_source == "saved_row_fold_gene_manifests" else "UNVERIFIED",
        (
            "Saved manifests establish the row assignment used for eligibility; "
            "reconstructed manifests are weaker evidence."
        ),
    )
    add(
        "Provenance",
        "Checkpoint-to-row-fold linkage",
        "Original training log/indices or exact prediction reproduction",
        "Not independently established by model and manifest filenames",
        "UNVERIFIED",
        (
            "The set-exclusion audit is valid, but the reviewer-grade historical "
            "claim remains conditional until each current checkpoint is linked "
            "to its saved row-fold manifest."
        ),
    )
    return pd.DataFrame(rows)


def make_gene_leakage_audit_figure(
    eligibility: Mapping[int, FoldEligibility],
    source_genes: np.ndarray,
    output_dir: Path,
    source: str,
    proof_source: str,
) -> Dict[str, Path]:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.colors import ListedColormap

    folds = sorted(eligibility)
    fold_sets = {
        fold: set(eligibility[fold].heldout_genes) for fold in folds
    }
    overlaps = np.zeros((len(folds), len(folds)), dtype=int)
    for row, left in enumerate(folds):
        for col, right in enumerate(folds):
            overlaps[row, col] = len(fold_sets[left] & fold_sets[right])

    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 9,
            "axes.spines.top": False,
            "axes.spines.right": False,
        }
    )
    fig, axes = plt.subplots(1, 2, figsize=(10.5, 4.2))
    fig.patch.set_facecolor("white")

    display = (overlaps > 0).astype(int)
    axes[0].imshow(display, cmap=ListedColormap(["#F8FAFC", "#FEF3C7"]), vmin=0, vmax=1)
    axes[0].set_xticks(range(len(folds)), [f"Fold {fold}" for fold in folds], rotation=35, ha="right")
    axes[0].set_yticks(range(len(folds)), [f"Fold {fold}" for fold in folds])
    axes[0].set_title("Unique genes shared by row folds")
    for row in range(len(folds)):
        for col in range(len(folds)):
            if row == col:
                label = f"{overlaps[row, col]}\ngenes"
            else:
                label = str(overlaps[row, col])
            axes[0].text(col, row, label, ha="center", va="center", fontsize=8)
    axes[0].text(
        0.5,
        -0.27,
        "Shared genes are excluded from the corresponding unseen-gene test",
        transform=axes[0].transAxes,
        ha="center",
        color="#92400E",
        fontweight="bold",
        fontsize=8,
    )

    x = np.arange(len(folds))
    eligible_counts = [len(eligibility[fold].eligible_genes) for fold in folds]
    excluded_counts = [len(eligibility[fold].excluded_seen_genes) for fold in folds]
    width = 0.36
    eligible_bars = axes[1].bar(
        x - width / 2,
        eligible_counts,
        width,
        color="#16A34A",
        label="Evaluated: absent from training",
    )
    excluded_bars = axes[1].bar(
        x + width / 2,
        excluded_counts,
        width,
        color="#F59E0B",
        label="Excluded: occurred in training",
    )
    axes[1].bar_label(eligible_bars, padding=2, fontsize=8)
    axes[1].bar_label(excluded_bars, padding=2, fontsize=8)
    axes[1].set_xticks(x, [f"Fold {fold}" for fold in folds])
    axes[1].set_ylabel("Number of unique genes")
    axes[1].set_title("Fold-specific target-gene eligibility")
    axes[1].grid(axis="y", color="#E2E8F0", linewidth=0.7)
    axes[1].legend(frameon=False, fontsize=7.5)
    axes[1].text(
        0.5,
        -0.22,
        "Training ∩ evaluated target genes = 0 in every fold",
        transform=axes[1].transAxes,
        ha="center",
        color="#166534",
        fontweight="bold",
        fontsize=8,
    )

    fig.suptitle(
        f"Fully unseen-gene audit: {source.replace('_', ' ')} source models",
        fontsize=12,
        fontweight="bold",
        y=1.02,
    )
    provenance_note = (
        "Set exclusion PASS; checkpoint-to-manifest linkage remains UNVERIFIED."
        if proof_source == "saved_row_fold_gene_manifests"
        else "Reconstructed assignment only; historical checkpoint linkage UNVERIFIED."
    )
    fig.text(0.5, -0.01, provenance_note, ha="center", fontsize=8.5, color="#9A3412")
    fig.tight_layout()

    stem = output_dir / f"{source}_gene_leakage_audit"
    png_path = stem.with_suffix(".png")
    pdf_path = stem.with_suffix(".pdf")
    fig.savefig(png_path, dpi=300, bbox_inches="tight", facecolor="white")
    fig.savefig(pdf_path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return {"gene_audit_png": png_path, "gene_audit_pdf": pdf_path}


def make_metrics_heatmap(
    summary: pd.DataFrame, output_dir: Path, source: str
) -> Dict[str, Path]:
    if summary.empty:
        return {}
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    columns = [
        "AUROC", "AUPRC", "accuracy", "precision", "recall", "specificity", "f1"
    ]
    matrix = summary.set_index("target")[columns].astype(float)
    labels = [index.replace("_", " ") for index in matrix.index]
    fig, axis = plt.subplots(figsize=(9.2, max(3.6, 0.55 * len(matrix) + 1.5)))
    image = axis.imshow(matrix.to_numpy(), cmap="YlGnBu", vmin=0.0, vmax=1.0, aspect="auto")
    display_columns = {
        "AUROC": "AUROC",
        "AUPRC": "AUPRC",
        "accuracy": "Accuracy",
        "precision": "Precision",
        "recall": "Recall",
        "specificity": "Specificity",
        "f1": "F1",
    }
    axis.set_xticks(
        range(len(columns)),
        [display_columns[column] for column in columns],
        rotation=35,
        ha="right",
    )
    axis.set_yticks(range(len(labels)), labels)
    axis.set_title("Cross-patient classification metrics", fontweight="bold", pad=12)
    for row in range(matrix.shape[0]):
        for col in range(matrix.shape[1]):
            value = matrix.iloc[row, col]
            color = "white" if value >= 0.65 else "#0F172A"
            axis.text(col, row, f"{value:.3f}", ha="center", va="center", color=color, fontsize=8)
    colorbar = fig.colorbar(image, ax=axis, fraction=0.025, pad=0.03)
    colorbar.set_label("Metric value")
    fig.tight_layout()
    stem = output_dir / f"{source}_cross_patient_metrics_heatmap"
    png_path = stem.with_suffix(".png")
    pdf_path = stem.with_suffix(".pdf")
    fig.savefig(png_path, dpi=300, bbox_inches="tight", facecolor="white")
    fig.savefig(pdf_path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return {"metrics_heatmap_png": png_path, "metrics_heatmap_pdf": pdf_path}


def make_confusion_matrix_figure(
    summary: pd.DataFrame, output_dir: Path, source: str, threshold: float
) -> Dict[str, Path]:
    if summary.empty:
        return {}
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    count = len(summary)
    columns = min(3, count)
    rows = int(np.ceil(count / columns))
    fig, axes = plt.subplots(rows, columns, figsize=(3.5 * columns, 3.2 * rows), squeeze=False)
    for axis in axes.ravel():
        axis.axis("off")
    for axis, (_, row) in zip(axes.ravel(), summary.iterrows()):
        axis.axis("on")
        matrix = np.asarray([[row["TN"], row["FP"]], [row["FN"], row["TP"]]], dtype=int)
        axis.imshow(matrix, cmap="Blues")
        axis.set_xticks([0, 1], ["NG", "CD"])
        axis.set_yticks([0, 1], ["NG", "CD"])
        axis.set_xlabel("Predicted")
        axis.set_ylabel("True")
        axis.set_title(str(row["target"]).replace("_", " "), fontweight="bold")
        for matrix_row in range(2):
            for matrix_col in range(2):
                axis.text(
                    matrix_col,
                    matrix_row,
                    f"{matrix[matrix_row, matrix_col]:,}",
                    ha="center",
                    va="center",
                    color="white" if matrix[matrix_row, matrix_col] > matrix.max() / 2 else "#0F172A",
                    fontweight="bold",
                )
    fig.suptitle(
        f"Cross-patient confusion matrices at threshold {threshold:g}",
        fontweight="bold",
        y=1.01,
    )
    fig.tight_layout()
    stem = output_dir / f"{source}_cross_patient_confusion_matrices"
    png_path = stem.with_suffix(".png")
    pdf_path = stem.with_suffix(".pdf")
    fig.savefig(png_path, dpi=300, bbox_inches="tight", facecolor="white")
    fig.savefig(pdf_path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return {"confusion_matrices_png": png_path, "confusion_matrices_pdf": pdf_path}


def write_excel_report(
    summary: pd.DataFrame,
    fold_summary: pd.DataFrame,
    predictions: pd.DataFrame,
    manifest: pd.DataFrame,
    leakage_audit: pd.DataFrame,
    report: Mapping[str, object],
    output_dir: Path,
    source: str,
    figure_paths: Mapping[str, Path],
) -> Path:
    """Write a styled, review-ready Excel workbook with raw and audit data."""
    try:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        from openpyxl.chart.series import SeriesLabel
        from openpyxl.formatting.rule import ColorScaleRule
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as error:
        raise AuditError(
            "Excel output requires openpyxl. Install it once with "
            "'conda install openpyxl' or run with --skip-excel."
        ) from error

    navy = "17365D"
    teal = "1F9E89"
    light_blue = "DDEBF7"
    light_green = "E2F0D9"
    light_orange = "FCE4D6"
    light_red = "F4CCCC"
    grey = "E7E6E6"
    white = "FFFFFF"
    thin_grey = Side(style="thin", color="D9E2F3")

    workbook = Workbook()
    dashboard = workbook.active
    dashboard.title = "Dashboard"
    metrics_sheet = workbook.create_sheet("Pooled Metrics")
    fold_metrics_sheet = workbook.create_sheet("Fold Metrics")
    prediction_sheet = workbook.create_sheet("Predictions")
    manifest_sheet = workbook.create_sheet("Fold Manifest")
    leakage_sheet = workbook.create_sheet("Gene Leakage Audit")
    readme_sheet = workbook.create_sheet("README")

    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False

    def excel_value(value: object) -> object:
        if isinstance(value, Path):
            return str(value)
        if isinstance(value, (np.integer, np.floating)):
            return value.item()
        if value is None:
            return None
        try:
            if pd.isna(value):
                return None
        except (TypeError, ValueError):
            pass
        if isinstance(value, (dict, list, tuple, set)):
            return json.dumps(value, default=JSON_value)
        return value

    def style_dataframe_sheet(
        sheet,
        frame: pd.DataFrame,
        table_name: str,
        freeze: str = "A2",
    ) -> None:
        headers = [str(column) for column in frame.columns]
        sheet.append(headers)
        for record in frame.itertuples(index=False, name=None):
            sheet.append([excel_value(value) for value in record])
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(color=white, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 26
        sheet.freeze_panes = freeze
        sheet.auto_filter.ref = sheet.dimensions
        if not frame.empty:
            table = Table(displayName=table_name, ref=sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            sheet.add_table(table)
        for column_cells in sheet.columns:
            letter = column_cells[0].column_letter
            maximum = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in column_cells[: min(len(column_cells), 200)]
            )
            sheet.column_dimensions[letter].width = min(max(maximum + 2, 10), 42)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=False)

    style_dataframe_sheet(metrics_sheet, summary, "MetricsTable")
    style_dataframe_sheet(fold_metrics_sheet, fold_summary, "FoldMetricsTable")
    style_dataframe_sheet(prediction_sheet, predictions, "PredictionsTable")
    style_dataframe_sheet(manifest_sheet, manifest, "FoldManifestTable")
    style_dataframe_sheet(leakage_sheet, leakage_audit, "GeneLeakageAuditTable")

    leakage_sheet.column_dimensions["A"].width = 18
    leakage_sheet.column_dimensions["B"].width = 48
    leakage_sheet.column_dimensions["C"].width = 36
    leakage_sheet.column_dimensions["D"].width = 55
    leakage_sheet.column_dimensions["E"].width = 15
    leakage_sheet.column_dimensions["F"].width = 76
    for row in leakage_sheet.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        row[2].alignment = Alignment(wrap_text=True, vertical="top")
        row[3].alignment = Alignment(wrap_text=True, vertical="top")
        row[5].alignment = Alignment(wrap_text=True, vertical="top")
        leakage_sheet.row_dimensions[row[0].row].height = 34

    for sheet in (metrics_sheet, fold_metrics_sheet, prediction_sheet):
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"

    if not summary.empty:
        metric_names = {str(cell.value): cell.column for cell in metrics_sheet[1]}
        metric_columns = [
            metric_names[name]
            for name in ("AUROC", "AUPRC", "accuracy", "precision", "recall", "specificity", "f1")
            if name in metric_names
        ]
        for column in metric_columns:
            letter = metrics_sheet.cell(row=1, column=column).column_letter
            metrics_sheet.conditional_formatting.add(
                f"{letter}2:{letter}{metrics_sheet.max_row}",
                ColorScaleRule(
                    start_type="num", start_value=0.0, start_color="F8696B",
                    mid_type="num", mid_value=0.5, mid_color="FFEB84",
                    end_type="num", end_value=1.0, end_color="63BE7B",
                ),
            )

    status_column = None
    for cell in leakage_sheet[1]:
        if cell.value == "status":
            status_column = cell.column
            break
    if status_column is not None:
        for row in range(2, leakage_sheet.max_row + 1):
            cell = leakage_sheet.cell(row=row, column=status_column)
            if cell.value in ("PASS", "VERIFIED"):
                cell.fill = PatternFill("solid", fgColor=light_green)
                cell.font = Font(color="166534", bold=True)
            elif cell.value == "UNVERIFIED":
                cell.fill = PatternFill("solid", fgColor=light_orange)
                cell.font = Font(color="9A3412", bold=True)
            elif cell.value == "FAIL":
                cell.fill = PatternFill("solid", fgColor=light_red)
                cell.font = Font(color="991B1B", bold=True)
            else:
                cell.fill = PatternFill("solid", fgColor=grey)

    dashboard.merge_cells("A1:H1")
    dashboard["A1"] = "OriGENE Cross-Patient Evaluation"
    dashboard["A1"].fill = PatternFill("solid", fgColor=navy)
    dashboard["A1"].font = Font(color=white, bold=True, size=16)
    dashboard["A1"].alignment = Alignment(horizontal="center", vertical="center")
    dashboard.row_dimensions[1].height = 34

    gene_checks = leakage_audit[leakage_audit.status.isin(["PASS", "FAIL"])]
    gene_partition_pass = not gene_checks.empty and bool((gene_checks.status == "PASS").all())
    provenance_verified = bool((leakage_audit.status == "VERIFIED").any())
    dashboard_rows = [
        ("Source patient", source.replace("_", " ")),
        ("Analysis status", str(report.get("status", ""))),
        ("Gene-partition audit", "PASS" if gene_partition_pass else "FAIL"),
        ("Historical checkpoint provenance", "VERIFIED" if provenance_verified else "UNVERIFIED"),
        ("Target-gene rule", "Only genes absent from that checkpoint's source-training folds"),
        ("Fold prediction averaging", "No — only an eligible checkpoint predicts each retained gene"),
        ("Held-out-gene source", str(report.get("heldout_gene_proof_source", ""))),
    ]
    for row_index, (label, value) in enumerate(dashboard_rows, start=3):
        dashboard.cell(row=row_index, column=1, value=label)
        dashboard.cell(row=row_index, column=2, value=value)
        dashboard.cell(row=row_index, column=1).font = Font(bold=True, color=navy)
        dashboard.cell(row=row_index, column=1).fill = PatternFill("solid", fgColor=light_blue)
        dashboard.cell(row=row_index, column=2).border = Border(bottom=thin_grey)
    dashboard["B5"].fill = PatternFill("solid", fgColor=light_green if gene_partition_pass else light_red)
    dashboard["B5"].font = Font(bold=True, color="166534" if gene_partition_pass else "991B1B")
    dashboard["B6"].fill = PatternFill("solid", fgColor=light_green if provenance_verified else light_orange)
    dashboard["B6"].font = Font(bold=True, color="166534" if provenance_verified else "9A3412")

    dashboard.merge_cells("A10:H11")
    dashboard["A10"] = (
        "Conclusion: only target genes absent from the corresponding source-training "
        "folds are evaluated. Historical checkpoint-to-manifest linkage still requires "
        "the original training records or exact reproduction of saved predictions."
    )
    dashboard["A10"].alignment = Alignment(wrap_text=True, vertical="center")
    dashboard["A10"].fill = PatternFill("solid", fgColor=light_orange if not provenance_verified else light_green)
    dashboard["A10"].font = Font(color="7C2D12" if not provenance_verified else "166534", italic=True)

    metric_start = 14
    if not summary.empty:
        dashboard.cell(row=metric_start, column=1, value="Performance summary")
        dashboard.cell(row=metric_start, column=1).font = Font(size=12, bold=True, color=navy)
        selected = summary[["target", "N", "AUROC", "AUPRC", "accuracy", "precision", "recall", "specificity", "f1"]]
        for column_index, column in enumerate(selected.columns, start=1):
            cell = dashboard.cell(row=metric_start + 1, column=column_index, value=column)
            cell.fill = PatternFill("solid", fgColor=teal)
            cell.font = Font(color=white, bold=True)
        for row_offset, record in enumerate(selected.itertuples(index=False, name=None), start=metric_start + 2):
            for column_index, value in enumerate(record, start=1):
                cell = dashboard.cell(row=row_offset, column=column_index, value=excel_value(value))
                if isinstance(cell.value, float):
                    cell.number_format = "0.000"
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "AUROC and AUPRC by target patient"
        chart.y_axis.title = "Metric value"
        chart.y_axis.scaling.min = 0
        chart.y_axis.scaling.max = 1
        chart.y_axis.numFmt = "0.0"
        chart.height = 7.5
        chart.width = 15
        data = Reference(
            dashboard,
            min_col=3,
            max_col=4,
            min_row=metric_start + 1,
            max_row=metric_start + 1 + len(selected),
        )
        categories = Reference(
            dashboard,
            min_col=1,
            min_row=metric_start + 2,
            max_row=metric_start + 1 + len(selected),
        )
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        if len(chart.series) >= 2:
            chart.series[0].title = SeriesLabel(v="AUROC")
            chart.series[0].graphicalProperties.solidFill = "4472C4"
            chart.series[1].title = SeriesLabel(v="AUPRC")
            chart.series[1].graphicalProperties.solidFill = "1F9E89"
        dashboard.add_chart(chart, "K3")

    dashboard.freeze_panes = "A3"
    dashboard.column_dimensions["A"].width = 35
    dashboard.column_dimensions["B"].width = 68
    for letter in "CDEFGHI":
        dashboard.column_dimensions[letter].width = 13

    readme_rows = [
        ("Purpose", "Cross-patient evaluation using existing OriGENE fold checkpoints without retraining."),
        ("Prediction rule", "Checkpoint k predicts only target-patient genes absent from the unique gene IDs in its other four source row folds."),
        ("Pooling rule", "Eligible fold-specific target predictions are concatenated. A checkpoint that trained on a gene never contributes a prediction for that gene."),
        ("Gene-exclusion PASS", "For every fold, source-training gene IDs and evaluated target-test gene IDs have zero intersection."),
        ("Critical limitation", "The workbook verifies set exclusion using the supplied row-fold manifests. It does not independently link the current checkpoint bytes to those historical manifests."),
        ("Required publication wording", "Cross-patient evaluation restricted to fully unseen genes; report the retained count and checkpoint-provenance limitation."),
        ("Metrics", "AUROC, AUPRC, accuracy, precision, recall, specificity and F1; classification metrics use the configured fixed threshold."),
        ("Figures", "; ".join(str(path) for path in figure_paths.values())),
        ("Audit JSON", str(output_dir / f"{source}_gene_disjoint_audit.json")),
    ]
    readme_sheet.append(["Item", "Description"])
    for label, description in readme_rows:
        readme_sheet.append([label, description])
    for cell in readme_sheet[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color=white, bold=True)
    readme_sheet.column_dimensions["A"].width = 30
    readme_sheet.column_dimensions["B"].width = 110
    for row in readme_sheet.iter_rows(min_row=2):
        row[0].font = Font(bold=True, color=navy)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        readme_sheet.row_dimensions[row[0].row].height = 42
    readme_sheet.freeze_panes = "A2"

    workbook.properties.title = "OriGENE cross-patient gene-disjoint evaluation"
    workbook.properties.subject = "Performance metrics and gene leakage audit"
    workbook.properties.creator = "OriGENE evaluation script"
    output_path = output_dir / f"{source}_cross_patient_report.xlsx"
    workbook.save(output_path)
    return output_path


def JSON_value(value: object) -> object:
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        return float(value)
    raise TypeError(f"Cannot JSON-serialize {type(value).__name__}")


def build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Reuse OriGENE fold checkpoints for gene-disjoint cross-patient "
            "prediction by letting checkpoint k evaluate only held-out genes "
            "that are absent from its other training folds."
        )
    )
    parser.add_argument("--source", default="Lung_I", choices=sorted(DATASET_SPECS))
    parser.add_argument(
        "--targets",
        nargs="+",
        default=None,
        choices=sorted(DATASET_SPECS),
        help="Target patients; defaults to every configured patient except source.",
    )
    parser.add_argument(
        "--sequence-root", default="Data/Final_sequences"
    )
    parser.add_argument(
        "--new-sequence-root", default="Data/Final_sequences/New"
    )
    parser.add_argument(
        "--model-dir", default="results/Model_evaluation/{source}"
    )
    parser.add_argument(
        "--model-template",
        default="OriGENE_{source}_40kbp_singletrack_{fold}.keras",
        help="Checkpoint filename/path template using {source} and {fold}.",
    )
    parser.add_argument(
        "--heldout-dir",
        default=None,
        help="Directory containing Fold{k}_genes.npy files; defaults to model-dir.",
    )
    parser.add_argument(
        "--heldout-template",
        default=None,
        help=(
            "Held-out gene filename/path template using {source} and {fold}. "
            "If omitted, strict unambiguous auto-detection is attempted."
        ),
    )
    parser.add_argument(
        "--output-dir",
        default=(
            "results/Model_evaluation/"
            "Cross_patient_gene_disjoint/{source}"
        ),
    )
    parser.add_argument("--n-folds", type=int, default=5)
    parser.add_argument("--random-state", type=int, default=4)
    parser.add_argument("--threshold", type=float, default=0.5)
    parser.add_argument("--n-bootstraps", type=int, default=1000)
    parser.add_argument("--bootstrap-seed", type=int, default=20260817)
    parser.add_argument(
        "--reconstruct-folds",
        action="store_true",
        help=(
            "Reconstruct held-out folds with StratifiedKFold only if saved gene "
            "files are missing. This is weaker evidence for legacy checkpoints."
        ),
    )
    parser.add_argument(
        "--force-reconstruct-folds",
        action="store_true",
        help=(
            "Ignore every saved held-out-gene file and reconstruct five "
            "source-only folds with StratifiedKFold. Use this for the current "
            "checkpoints when existing manifests belong to another experiment. "
            "The output is explicitly marked as unproven checkpoint provenance."
        ),
    )
    parser.add_argument(
        "--audit-only",
        action="store_true",
        help="Run all file/gene/label/checkpoint audits without model prediction.",
    )
    parser.add_argument(
        "--skip-excel",
        action="store_true",
        help="Skip the styled Excel workbook; CSV, JSON and figure outputs remain.",
    )
    return parser


def run(args: argparse.Namespace) -> int:
    if args.n_folds < 2:
        raise AuditError("--n-folds must be at least 2.")
    if not 0.0 < args.threshold < 1.0:
        raise AuditError("--threshold must be strictly between 0 and 1.")
    targets = list(args.targets) if args.targets else [
        patient for patient in DATASET_SPECS if patient != args.source
    ]
    if args.source in targets:
        raise AuditError(
            f"Source patient {args.source} is also listed as a target; "
            "self-target evaluation is forbidden."
        )

    sequence_root = expand(args.sequence_root)
    new_root = expand(args.new_sequence_root)
    model_dir = expand(format_template(args.model_dir, args.source, 1))
    heldout_dir = expand(args.heldout_dir) if args.heldout_dir else model_dir
    output_dir = expand(format_template(args.output_dir, args.source, 1))
    output_dir.mkdir(parents=True, exist_ok=True)

    source_labels, source_genes = load_patient_metadata(
        args.source, sequence_root, new_root
    )
    checkpoint_paths = resolve_model_paths(
        model_dir, args.model_template, args.source, args.n_folds
    )
    checkpoint_hashes = {
        fold: sha256_file(path) for fold, path in checkpoint_paths.items()
    }

    fold_rows, heldout_paths, proof_source = load_heldout_folds(
        heldout_dir=heldout_dir,
        template=args.heldout_template,
        source=args.source,
        source_genes=source_genes,
        source_labels=source_labels,
        n_folds=args.n_folds,
        reconstruct=args.reconstruct_folds,
        force_reconstruct=args.force_reconstruct_folds,
        random_state=args.random_state,
    )
    eligibility, fold_audit = derive_fold_eligibility(fold_rows, source_genes)

    source_label_map = dict(zip(source_genes, source_labels))
    manifest_rows: List[Dict[str, object]] = []
    for fold, details in eligibility.items():
        checkpoint = checkpoint_paths[fold]
        heldout_counts = Counter(details.heldout_rows.tolist())
        training_row_counts = Counter(
            gene
            for other_fold, rows in fold_rows.items()
            if other_fold != fold
            for gene in rows.tolist()
        )
        eligible_set = set(details.eligible_genes)
        for gene in details.heldout_genes:
            eligible = gene in eligible_set
            manifest_rows.append(
                {
                    "source": args.source,
                    "fold": fold,
                    "gene": gene,
                    "label": int(source_label_map[gene]),
                    "heldout_row_occurrences": int(heldout_counts[gene]),
                    "source_training_row_occurrences": int(training_row_counts[gene]),
                    "source_training_gene_present": bool(training_row_counts[gene] > 0),
                    "eligible_fully_unseen": bool(eligible),
                    "target_role_for_checkpoint": (
                        "EVALUATED" if eligible else "EXCLUDED_SEEN_IN_TRAINING"
                    ),
                    "checkpoint": str(checkpoint),
                    "checkpoint_sha256": checkpoint_hashes[fold],
                    "row_fold_manifest_source": proof_source,
                    "heldout_gene_file": str(heldout_paths.get(fold, "")),
                }
            )
    manifest = pd.DataFrame(manifest_rows).sort_values(["fold", "gene"])
    manifest_path = output_dir / f"{args.source}_fold_gene_manifest.csv"
    manifest.to_csv(manifest_path, index=False)
    eligible_union = set().union(
        *(set(details.eligible_genes) for details in eligibility.values())
    )

    report: Dict[str, object] = {
        "status": "AUDIT_IN_PROGRESS",
        "created_utc": datetime.now(timezone.utc).isoformat(),
        "source": args.source,
        "targets": targets,
        "sequence_root": sequence_root,
        "new_sequence_root": new_root,
        "model_dir": model_dir,
        "heldout_dir": heldout_dir,
        "output_dir": output_dir,
        "n_folds": args.n_folds,
        "fold_random_state": args.random_state,
        "heldout_gene_proof_source": proof_source,
        "checkpoint_paths": {
            str(fold): path for fold, path in checkpoint_paths.items()
        },
        "checkpoint_sha256": {
            str(fold): digest for fold, digest in checkpoint_hashes.items()
        },
        "heldout_gene_paths": {
            str(fold): path for fold, path in heldout_paths.items()
        },
        "source_gene_order_sha256": sha256_strings(source_genes),
        "source_sorted_gene_set_sha256": sha256_strings(sorted(source_genes)),
        "fold_audit": fold_audit,
        "target_audits": {},
        "manifest_csv": manifest_path,
        "limitations": [],
    }
    report["limitations"].append(
        "Confirm from the original training code that only the declared source "
        "patient entered model fitting and that outer held-out genes did not "
        "control early stopping, checkpoint selection, preprocessing, feature "
        "selection, or hyperparameter tuning."
    )
    report["limitations"].append(
        "The supplied row-fold manifests identify which genes are absent from "
        "each fold's other four row sets. They do not, by themselves, prove that "
        "the current checkpoint bytes were trained from those exact rows."
    )
    if "unproven_checkpoint_fold_provenance" in proof_source:
        report["limitations"].append(
            "Fold membership was reconstructed rather than loaded from saved "
            "per-fold gene files; legacy checkpoint provenance and checkpoint-to-"
            "fold numbering are not independently proven. This analysis prevents "
            "evaluation-time fold mixing but cannot rule out training-time leakage."
        )

    target_metadata: Dict[str, Tuple[np.ndarray, np.ndarray]] = {}
    for target in targets:
        target_labels, target_genes = load_patient_metadata(
            target, sequence_root, new_root
        )
        target_metadata[target] = (target_labels, target_genes)
        report["target_audits"][target] = target_gene_audit(
            source_genes,
            source_labels,
            target_genes,
            target_labels,
            args.source,
            target,
        )
        report["target_audits"][target]["eligible_gene_count"] = len(
            eligible_union
        )

    audit_path = output_dir / f"{args.source}_gene_disjoint_audit.json"
    if args.audit_only:
        report["status"] = (
            "FULLY_UNSEEN_GENE_SUBSET_AUDIT_PASS_"
            "CHECKPOINT_LINK_UNVERIFIED"
        )
        leakage_audit = build_gene_leakage_audit_table(
            eligibility,
            source_genes,
            report["target_audits"],
            proof_source,
        )
        leakage_csv = output_dir / f"{args.source}_gene_leakage_audit.csv"
        leakage_audit.to_csv(leakage_csv, index=False)
        figure_paths = make_gene_leakage_audit_figure(
            eligibility,
            source_genes,
            output_dir,
            args.source,
            proof_source,
        )
        report["gene_partition_audit_csv"] = leakage_csv
        report["figure_paths"] = figure_paths
        if not args.skip_excel:
            excel_path = write_excel_report(
                summary=pd.DataFrame(),
                fold_summary=pd.DataFrame(),
                predictions=pd.DataFrame(),
                manifest=manifest,
                leakage_audit=leakage_audit,
                report=report,
                output_dir=output_dir,
                source=args.source,
                figure_paths=figure_paths,
            )
            report["excel_report"] = excel_path
        with audit_path.open("w", encoding="utf-8") as handle:
            json.dump(report, handle, indent=2, default=JSON_value)
        print(report["status"])
        print(f"Manifest: {manifest_path}")
        print(f"Gene leakage audit: {leakage_csv}")
        if "excel_report" in report:
            print(f"Excel report: {report['excel_report']}")
        print(f"Gene audit figure: {figure_paths['gene_audit_png']}")
        print(f"Audit: {audit_path}")
        return 0

    all_predictions: List[pd.DataFrame] = []
    summary_rows: List[Dict[str, object]] = []
    fold_summary_rows: List[Dict[str, object]] = []

    for target_index, target in enumerate(targets):
        print(f"\n=== {args.source} -> {target}: gene-disjoint prediction ===")
        x_target, y_target, g_target = load_patient_gene_subset(
            target,
            sorted(eligible_union),
            sequence_root,
            new_root,
        )
        subset_audit = evaluated_target_subset_audit(
            source_genes,
            source_labels,
            g_target,
            y_target,
            eligible_union,
            args.source,
            target,
        )
        report["target_audits"][target].update(subset_audit)
        target_position = index_by_gene(g_target, target)

        fold_predictions: List[pd.DataFrame] = []
        for fold in range(1, args.n_folds + 1):
            details = eligibility[fold]
            test_genes = details.eligible_genes
            test_set = set(test_genes)
            source_training_set = set(details.training_genes)
            if not source_training_set.isdisjoint(test_set):
                raise AuditError(
                    f"Fold {fold}: source-training/target-test overlap is non-zero."
                )
            if len(test_genes) == 0:
                raise AuditError(
                    f"Fold {fold}: no fully unseen genes remain after excluding "
                    "genes present in the other four training folds."
                )

            target_indices = np.asarray(
                [target_position[gene] for gene in test_genes], dtype=int
            )
            x_test = x_target[target_indices]
            y_test = y_target[target_indices]

            checkpoint = checkpoint_paths[fold]
            print(
                f"Fold {fold}: checkpoint={checkpoint.name}; "
                f"training unique genes={len(source_training_set)}; "
                f"eligible target genes={len(test_genes)}; "
                f"excluded seen genes={len(details.excluded_seen_genes)}"
            )
            model = load_epi_model(checkpoint)
            validate_model_input(model, x_test, checkpoint)
            scores = np.asarray(model.predict(x_test, verbose=0)).reshape(-1)
            if len(scores) != len(test_genes):
                raise AuditError(
                    f"Fold {fold}: model returned {len(scores)} predictions for "
                    f"{len(test_genes)} genes."
                )
            if not np.all(np.isfinite(scores)):
                raise AuditError(f"Fold {fold}: predictions contain NaN or infinity.")
            clear_model(model)
            del model

            fold_frame = pd.DataFrame(
                {
                    "source": args.source,
                    "target": target,
                    "fold": fold,
                    "gene": test_genes,
                    "label": y_test.astype(int),
                    "score": scores.astype(float),
                    "source_training_gene_present": False,
                    "gene_unseen_by_declared_training_folds": True,
                    "checkpoint": str(checkpoint),
                    "checkpoint_sha256": checkpoint_hashes[fold],
                }
            )
            fold_predictions.append(fold_frame)

            fold_metrics = classification_metrics(
                fold_frame.label.to_numpy(dtype=int),
                fold_frame.score.to_numpy(dtype=float),
                args.threshold,
            )
            fold_metrics.update(
                {
                    "source": args.source,
                    "target": target,
                    "fold": fold,
                    "source_training_unique_genes": len(source_training_set),
                    "saved_heldout_rows": len(details.heldout_rows),
                    "saved_heldout_unique_genes": len(details.heldout_genes),
                    "eligible_fully_unseen_genes": len(details.eligible_genes),
                    "excluded_seen_genes": len(details.excluded_seen_genes),
                    "source_training_vs_evaluated_overlap": 0,
                    "checkpoint": str(checkpoint),
                    "checkpoint_sha256": checkpoint_hashes[fold],
                }
            )
            fold_summary_rows.append(fold_metrics)

        pooled = pd.concat(fold_predictions, ignore_index=True)
        if pooled.gene.duplicated().any():
            duplicated = pooled.loc[pooled.gene.duplicated(), "gene"].unique()
            raise AuditError(
                f"{target}: genes were predicted more than once; examples: "
                f"{duplicated[:10].tolist()}. No fold averaging is allowed."
            )
        if set(pooled.gene) != eligible_union or len(pooled) != len(eligible_union):
            raise AuditError(
                f"{target}: pooled predictions do not cover every eligible "
                "fully unseen gene exactly once."
            )

        pooled = pooled.sort_values("gene").reset_index(drop=True)
        metrics = classification_metrics(
            pooled.label.to_numpy(dtype=int),
            pooled.score.to_numpy(dtype=float),
            args.threshold,
        )
        metrics.update(
            bootstrap_intervals(
                pooled.label.to_numpy(dtype=int),
                pooled.score.to_numpy(dtype=float),
                args.n_bootstraps,
                args.bootstrap_seed + target_index,
            )
        )
        metrics.update(
            {
                "source": args.source,
                "target": target,
                "evaluation": "pooled_fully_unseen_gene_cross_patient_predictions",
                "source_unique_gene_universe": len(source_genes),
                "eligible_gene_count": len(eligible_union),
                "eligible_gene_fraction": len(eligible_union) / len(source_genes),
                "fold_prediction_averaging": False,
                "each_gene_predicted_once": True,
            }
        )

        prediction_path = (
            output_dir / f"{args.source}_to_{target}_gene_disjoint_predictions.csv"
        )
        pooled.to_csv(prediction_path, index=False)
        metrics["prediction_csv"] = str(prediction_path)
        summary_rows.append(metrics)
        all_predictions.append(pooled)
        report["target_audits"][target].update(
            {
                "pooled_prediction_count": len(pooled),
                "unique_predicted_gene_count": pooled.gene.nunique(),
                "each_gene_predicted_once": True,
                "fold_prediction_averaging": False,
                "prediction_csv": prediction_path,
            }
        )
        print(
            f"PASS {args.source} -> {target}: AUROC={metrics['AUROC']:.4f}, "
            f"AUPRC={metrics['AUPRC']:.4f}, N={metrics['N']}"
        )

        del x_target, y_target, g_target
        gc.collect()

    summary = pd.DataFrame(summary_rows).sort_values("target")
    fold_summary = pd.DataFrame(fold_summary_rows).sort_values(["target", "fold"])
    summary_path = output_dir / f"{args.source}_gene_disjoint_metrics.csv"
    fold_summary_path = output_dir / f"{args.source}_gene_disjoint_fold_metrics.csv"
    summary.to_csv(summary_path, index=False)
    fold_summary.to_csv(fold_summary_path, index=False)
    combined_predictions = pd.concat(all_predictions, ignore_index=True)
    combined_path = output_dir / f"{args.source}_gene_disjoint_all_predictions.csv"
    combined_predictions.to_csv(combined_path, index=False)

    report["status"] = "FULLY_UNSEEN_GENE_SUBSET_PASS_CHECKPOINT_LINK_UNVERIFIED"
    report["metrics_csv"] = summary_path
    report["fold_metrics_csv"] = fold_summary_path
    report["combined_predictions_csv"] = combined_path
    leakage_audit = build_gene_leakage_audit_table(
        eligibility,
        source_genes,
        report["target_audits"],
        proof_source,
    )
    leakage_csv = output_dir / f"{args.source}_gene_leakage_audit.csv"
    leakage_audit.to_csv(leakage_csv, index=False)
    figure_paths: Dict[str, Path] = {}
    figure_paths.update(make_plot(combined_predictions, output_dir, args.source))
    figure_paths.update(
        make_gene_leakage_audit_figure(
            eligibility,
            source_genes,
            output_dir,
            args.source,
            proof_source,
        )
    )
    figure_paths.update(make_metrics_heatmap(summary, output_dir, args.source))
    figure_paths.update(
        make_confusion_matrix_figure(
            summary, output_dir, args.source, args.threshold
        )
    )
    report["gene_partition_audit_csv"] = leakage_csv
    report["figure_paths"] = figure_paths
    if not args.skip_excel:
        excel_path = write_excel_report(
            summary=summary,
            fold_summary=fold_summary,
            predictions=combined_predictions,
            manifest=manifest,
            leakage_audit=leakage_audit,
            report=report,
            output_dir=output_dir,
            source=args.source,
            figure_paths=figure_paths,
        )
        report["excel_report"] = excel_path
    with audit_path.open("w", encoding="utf-8") as handle:
        json.dump(report, handle, indent=2, default=JSON_value)

    print(f"\n{report['status']}")
    print(f"Metrics: {summary_path}")
    print(f"Fold metrics: {fold_summary_path}")
    print(f"Predictions: {combined_path}")
    print(f"Manifest: {manifest_path}")
    print(f"Gene leakage audit: {leakage_csv}")
    if "excel_report" in report:
        print(f"Excel report: {report['excel_report']}")
    print(f"ROC/PR figure: {figure_paths['roc_pr_png']}")
    print(f"Gene audit figure: {figure_paths['gene_audit_png']}")
    print(f"Metrics heatmap: {figure_paths['metrics_heatmap_png']}")
    print(f"Confusion matrices: {figure_paths['confusion_matrices_png']}")
    print(f"Audit: {audit_path}")
    return 0


def main() -> int:
    parser = build_argument_parser()
    args = parser.parse_args()
    try:
        return run(args)
    except (AuditError, FileNotFoundError, RuntimeError, ValueError) as error:
        print(f"AUDIT_FAIL: {error}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
